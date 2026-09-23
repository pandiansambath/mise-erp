"""The consolidated Monthly Business Report.

    "i need one consolidated super spceial export feature in pnl area that
     includes litrelly al thre details like expense , sales, money ectetc (for
     particluar days or whatever we choose)"
    "have in excel, csv, pdf, words in all fomr"
    "ask user whther to iunclude tis calcuateion or not or leave empty to keep
     default whihc give all the consoldiated things as our sampe report"

Modelled on `docs/Nirai_August_2026_Updated_MBR_v2.pdf`, which he sent as the
spec. What these hold is the two things that make the report trustworthy: the
picker's EMPTY case means everything, and a figure we cannot work out prints as
a dash rather than as a confident zero.
"""
from decimal import Decimal

import pytest

from app.auth.models import Role
from app.reports import mbr, mbr_export


@pytest.fixture
async def owner(make_user):
    return await make_user("owner@mbr.test", Role.SUPER_ADMIN.value)


# ── the picker ────────────────────────────────────────────────────────────


async def test_the_catalogue_is_offered(client, auth_header, owner) -> None:
    """The chooser is SERVED, not hardcoded in the browser — a section added to
    `mbr.py` has to appear in the picker without a second edit."""
    res = await client.get("/api/reports/mbr/sections", headers=auth_header(owner))
    assert res.status_code == 200
    keys = [s["key"] for s in res.json()["sections"]]
    assert keys == mbr.ALL_KEYS
    assert all(s["label"] for s in res.json()["sections"]), "every section needs a label"


async def test_asking_for_nothing_gives_everything(client, auth_header, owner) -> None:
    """⭐ HIS RULE: "leave empty to keep default whihc give all the consoldiated
    things as our sampe report". Empty is how somebody says "the usual", not
    how they ask for a blank document."""
    res = await client.get(
        "/api/reports/mbr?date_from=2026-08-01&date_to=2026-08-31",
        headers=auth_header(owner),
    )
    assert res.status_code == 200
    assert [s["key"] for s in res.json()["sections"]] == mbr.ALL_KEYS


async def test_asking_for_two_gives_two(client, auth_header, owner) -> None:
    res = await client.get(
        "/api/reports/mbr?date_from=2026-08-01&date_to=2026-08-31&sections=summary,vendors",
        headers=auth_header(owner),
    )
    assert res.status_code == 200
    assert [s["key"] for s in res.json()["sections"]] == ["summary", "vendors"]


async def test_a_made_up_section_is_ignored_not_fatal(client, auth_header, owner) -> None:
    """An unknown key is dropped rather than 500ing. A stale bookmark should
    still produce a report."""
    res = await client.get(
        "/api/reports/mbr?date_from=2026-08-01&date_to=2026-08-31&sections=summary,nonsense",
        headers=auth_header(owner),
    )
    assert res.status_code == 200
    assert [s["key"] for s in res.json()["sections"]] == ["summary"]


# ── the four downloads ────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("fmt", "magic"),
    [("csv", b"\xef\xbb\xbf"), ("xlsx", b"PK"), ("pdf", b"%PDF"), ("docx", b"PK")],
)
async def test_every_format_downloads(client, auth_header, owner, fmt, magic) -> None:
    """⭐ "excel, csv, pdf, words in all fomr".

    The magic bytes are checked, not just the status: a 200 carrying an empty
    or mislabelled body fails on his machine rather than here, which is the
    worst place to find out.
    """
    res = await client.get(
        f"/api/reports/mbr.{fmt}?date_from=2026-08-01&date_to=2026-08-31",
        headers=auth_header(owner),
    )
    assert res.status_code == 200, res.text
    assert res.content[: len(magic)] == magic, f"{fmt} is not really a {fmt}"
    assert "attachment;" in res.headers["content-disposition"]


async def test_an_unknown_format_is_refused_in_words(client, auth_header, owner) -> None:
    res = await client.get(
        "/api/reports/mbr.rtf?date_from=2026-08-01&date_to=2026-08-31",
        headers=auth_header(owner),
    )
    assert res.status_code == 404
    assert "csv" in res.json()["detail"]


async def test_the_report_needs_permission(client, auth_header, make_user) -> None:
    staff = await make_user("staff@mbr.test", Role.STAFF.value)
    res = await client.get(
        "/api/reports/mbr.pdf?date_from=2026-08-01&date_to=2026-08-31",
        headers=auth_header(staff),
    )
    assert res.status_code == 403


# ── the thing that makes it trustworthy ───────────────────────────────────


def test_a_figure_we_cannot_work_out_is_a_dash() -> None:
    """⭐ NOT 0.00, EVER.

    His own report writes "Not supplied" and "Not calculated" in three places,
    and that honesty is the most valuable thing in it. A percentage of nothing
    is unknown, not zero, and printing the two the same way is how a report
    starts lying quietly.
    """
    assert mbr_export.fmt(None, money=True) == "—"
    assert mbr_export.fmt(None) == "—"
    assert mbr_export.fmt(Decimal("0"), money=True) == "£0.00"
    assert mbr._pct(Decimal("5"), Decimal("0")) is None


def test_a_loss_puts_the_sign_before_the_symbol() -> None:
    """Found on the LIVE report, not in review: a bad month printed as
    "£-1,465.01", which reads as a typo. Every statement a restaurant has ever
    seen writes it the other way round."""
    assert mbr_export.fmt(Decimal("-1465.01"), money=True) == "-£1,465.01"
    assert mbr_export.fmt(Decimal("1465.01"), money=True) == "£1,465.01"


def test_money_is_formatted_in_the_hotels_currency() -> None:
    assert mbr_export.fmt(Decimal("1234.5"), money=True, currency="GBP") == "£1,234.50"
    assert mbr_export.fmt(Decimal("1234.5"), money=True, currency="INR") == "₹1,234.50"


def test_excel_holds_numbers_not_text() -> None:
    """Most of why somebody wants Excel rather than the PDF is to sum a column,
    and "£1,234.00" as text cannot be summed."""
    from openpyxl import load_workbook

    report = {
        "hotel_name": "T", "date_from": "2026-08-01", "date_to": "2026-08-31",
        "currency": "GBP",
        "sections": [
            mbr.Section(
                key="vendors", title="Supplier payments",
                columns=["Supplier", "Paid"],
                rows=[["Meat Wala", Decimal("1996.56")]],
                total=["TOTAL", Decimal("1996.56")],
                money_cols=[1],
            )
        ],
    }
    import io as _io

    wb = load_workbook(_io.BytesIO(mbr_export.to_xlsx(report)))
    ws = wb["Supplier payments"]
    found = [
        c.value
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, int | float)
    ]
    assert 1996.56 in found, "the amount went in as text, so nobody can sum it"
