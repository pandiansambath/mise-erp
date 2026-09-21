"""Export and import defined ONCE, so a file we write is a file we can read.

    "add rule: whatever we export we can import and use the same"

His rule, made structural. The inventory round trip broke because the exporter
and the importer each kept their own list of English column names in their own
file, and the two drifted: the exporter wrote "In stock", the importer accepted
"Opening stock" and four aliases that did not include it, and every quantity was
dropped on re-import WITHOUT AN ERROR.

Fixing that one pair of lists would have left the same trap set in vendors, in
employees and in recipes. So the shape here is: ONE declaration per list, from
which both the CSV, the XLSX and the import template are generated. There is no
second place to keep in step, because there is no second list.

WHAT A ROUND TRIP ACTUALLY REQUIRES, learned the hard way
--------------------------------------------------------------------------
1. THE HEADERS MUST MATCH. Not "be similar" — match, by construction.
2. THE DECORATION MUST COME BACK OFF. A human export marks the chosen supplier
   "★ Fresh Farms" and writes "yes" for a boolean. Both are right for a person
   reading it and neither survives a naive re-read, so every field declares how
   to write itself AND how to read itself back.
3. A FOOTER IS NOT A ROW. Our exports end in totals. An importer that treats
   that as data fails the whole file on it — which is exactly what happened.
4. AN EMPTY EXPORT IS STILL A VALID FILE. A new restaurant exports nothing and
   must still get headers, so the file can be filled in by hand and imported.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Callable
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

from app.core.template_io import Column, TemplateSpec


@dataclass
class Field:
    """One column, in both directions.

    `key` is the name the importer hands back and the exporter reads off the
    row object. `header` is the one piece of English, written on the way out and
    accepted on the way in — it cannot drift because there is only one of it.
    """

    key: str
    header: str
    required: bool = False
    kind: str = "text"  # text | number | date
    #: Other spellings a human (or another system's export) might use.
    aliases: tuple[str, ...] = ()
    #: How to render this field for a person. Return "" for absent.
    to_cell: Callable[[Any], Any] | None = None
    #: How to read a person's cell back. Return None to reject the value.
    from_cell: Callable[[str], Any] | None = None
    #: Column width in the spreadsheet.
    width: int = 18
    #: Right-align (numbers and money).
    right: bool = False


@dataclass
class ListSpec:
    """A list that can leave and come back."""

    name: str
    title: str
    subtitle: str
    fields: list[Field]
    #: Shown in the downloadable blank template so an empty restaurant has an
    #: example to follow rather than a bare header row.
    sample_rows: list[list[Any]] = field(default_factory=list)

    def template(self) -> TemplateSpec:
        """The import spec, generated — never hand-written alongside this.

        Each column accepts its own export header, so the file we produce is by
        definition a file we accept.
        """
        return TemplateSpec(
            name=self.name,
            subtitle=self.subtitle,
            columns=[
                Column(
                    f.key,
                    f.header,
                    required=f.required,
                    kind=f.kind,
                    aliases=tuple(dict.fromkeys(a.lower() for a in f.aliases)),
                )
                for f in self.fields
            ],
            sample_rows=self.sample_rows,
        )

    def headers(self) -> list[str]:
        return [f.header for f in self.fields]

    def row_of(self, obj: Any) -> list[Any]:
        out: list[Any] = []
        for f in self.fields:
            raw = obj.get(f.key) if isinstance(obj, dict) else getattr(obj, f.key, None)
            out.append(f.to_cell(raw) if f.to_cell else _plain(raw))
        return out

    def clean(self, rec: dict) -> dict:
        """Undo the export's human decoration on the way back in."""
        out = dict(rec)
        for f in self.fields:
            if f.from_cell and f.key in out and isinstance(out[f.key], str):
                out[f.key] = f.from_cell(out[f.key])
        return out


def _plain(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, bool):
        # "yes"/"" reads better than True/False to a person, and `yes_no_back`
        # is the matching reader.
        return "yes" if v else ""
    return v


# ── the two readers every list needs ──────────────────────────────────────


def yes_no_back(s: str) -> bool:
    return s.strip().casefold() in {"yes", "y", "true", "1", "✓"}


def strip_mark(mark: str) -> Callable[[str], str]:
    """Drop a leading decoration such as the chosen-supplier star."""
    m = mark.strip()

    def _read(s: str) -> str:
        out = s.strip()
        while m and out.startswith(m):
            out = out[len(m) :].strip()
        return out

    return _read


# ── writing ───────────────────────────────────────────────────────────────


def to_csv(spec: ListSpec, rows: list[Any], *, footer: list[Any] | None = None) -> bytes:
    """A CSV a person can read and this module can read back.

    The title and blank line are for the person. The importer finds the header
    row by scanning, so they cost nothing — and a footer is skipped on re-read
    because it carries none of the required fields.
    """
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([spec.title])
    w.writerow([])
    w.writerow(spec.headers())
    for r in rows:
        w.writerow(spec.row_of(r))
    if footer:
        w.writerow([])
        w.writerow(footer)
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(spec: ListSpec, rows: list[Any], *, footer: list[Any] | None = None) -> bytes:
    from openpyxl import Workbook

    from app.core.xlsx_style import style_table

    wb = Workbook()
    ws = wb.active
    ws.title = spec.name[:31]
    for i, r in enumerate(rows):
        for c, val in enumerate(spec.row_of(r), start=1):
            ws.cell(row=4 + i, column=c, value=val)
    if footer:
        for c, val in enumerate(footer, start=1):
            ws.cell(row=4 + len(rows) + 1, column=c, value=val)
    style_table(
        ws,
        title=spec.title,
        subtitle=spec.subtitle,
        headers=spec.headers(),
        n_rows=max(len(rows), 1),
        widths=[f.width for f in spec.fields],
        right_cols={i for i, f in enumerate(spec.fields, start=1) if f.right},
    )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
