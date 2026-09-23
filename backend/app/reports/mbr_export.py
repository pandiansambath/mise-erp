"""The Monthly Business Report, rendered four ways.

    "have in excel, csv, pdf, words in all fomr we need to allwo user to
     dnwload... please use cool colors and decorotaew with stuff"

FOUR RENDERERS, ONE STRUCTURE. Each walks the same list of `Section` objects
from `mbr.py`, so a new section appears in all four documents at once. The
alternative — four functions that each know what a vendor table looks like — is
how three of them end up a month out of date with the fourth.

THE PALETTE IS SHARED, so the Excel workbook, the PDF and the Word document are
recognisably the same report rather than three things about the same numbers.
Teal for the headings, a warm amber for totals, and a stripe on alternate rows
because a wide money table is unreadable without one.

⚠️ CSV GETS NO DECORATION, ON PURPOSE. It is the one format here that is
machine input as often as it is a document, so it stays flat and quoted — a CSV
with merged title banners in it is a CSV nothing can parse.
"""
from __future__ import annotations

import csv
import io
from decimal import Decimal

from app.reports.mbr import Section

# ── the palette, shared by every renderer ─────────────────────────────────
INK = "0F3D3E"        # deep teal, headings
BRAND = "0E7C7B"      # teal, section bars
BRAND_SOFT = "D6EEEE"  # tint behind column headers
TOTAL_BG = "FDF1D6"   # warm amber, total rows
STRIPE = "F6F9F9"     # alternate rows
MUTED = "6B7280"

_RGB = {
    "INK": (15, 61, 62),
    "BRAND": (14, 124, 123),
    "BRAND_SOFT": (214, 238, 238),
    "TOTAL_BG": (253, 241, 214),
    "STRIPE": (246, 249, 249),
    "MUTED": (107, 114, 128),
}

_SYMBOL = {"GBP": "£", "USD": "$", "EUR": "€", "INR": "₹"}


def _is_money(section: Section, col: int) -> bool:
    return col in section.money_cols


def fmt(value, *, money: bool = False, currency: str = "GBP") -> str:
    """One cell, as text.

    ⚠️ `None` IS AN EM DASH, NEVER 0.00. A figure we could not work out and a
    figure that is genuinely nothing are different answers, and printing them
    the same way is how a report starts lying quietly.
    """
    if value is None:
        return "—"
    if isinstance(value, Decimal) and money:
        sym = _SYMBOL.get(currency, "")
        return f"{sym}{value:,.2f}"
    if isinstance(value, Decimal):
        return f"{value:,.2f}"
    return str(value)


def _title(report: dict) -> str:
    who = report.get("hotel_name") or "Your restaurant"
    return f"{who} — Business Report"


def _period(report: dict) -> str:
    return f"{report['date_from']} to {report['date_to']}"


# ══ CSV ═══════════════════════════════════════════════════════════════════


def to_csv(report: dict) -> bytes:
    """Flat, quoted, parseable. See the note at the top about decoration."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([_title(report)])
    w.writerow(["Period", _period(report)])
    cur = report["currency"]

    for s in report["sections"]:
        w.writerow([])
        w.writerow([s.title])
        if s.note:
            w.writerow([s.note])
        w.writerow(s.columns)
        for row in s.rows:
            w.writerow([fmt(v, money=_is_money(s, i), currency=cur) for i, v in enumerate(row)])
        if s.total:
            w.writerow([fmt(v, money=_is_money(s, i), currency=cur) for i, v in enumerate(s.total)])
    return buf.getvalue().encode("utf-8-sig")


# ══ Excel ═════════════════════════════════════════════════════════════════


def to_xlsx(report: dict) -> bytes:
    """One worksheet per section, plus a contents page.

    A single sheet with seven stacked tables is what people ask for and then
    cannot use: you cannot sort or filter one table without wrecking the one
    below it. A sheet each keeps every table sortable.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    cur = report["currency"]
    sym = _SYMBOL.get(cur, "")
    money_fmt = f'{sym}#,##0.00;[Red]-{sym}#,##0.00'

    wb = Workbook()
    contents = wb.active
    contents.title = "Contents"
    contents["A1"] = _title(report)
    contents["A1"].font = Font(bold=True, size=16, color=INK)
    contents["A2"] = f"Period: {_period(report)}"
    contents["A2"].font = Font(size=11, italic=True, color=MUTED)
    contents["A4"] = "What is in this report"
    contents["A4"].font = Font(bold=True, size=12, color=BRAND)
    contents.column_dimensions["A"].width = 46
    contents.column_dimensions["B"].width = 16

    thin = Side(style="thin", color="E3EAEA")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    for n, s in enumerate(report["sections"], start=5):
        contents.cell(row=n, column=1, value=s.title).font = Font(color=INK)
        contents.cell(row=n, column=2, value=f"{len(s.rows)} rows").font = Font(color=MUTED)

        # Excel sheet titles: 31 chars, and none of : \ / ? * [ ]
        safe = "".join(c for c in s.title if c not in ':\\/?*[]')[:31]
        ws = wb.create_sheet(safe or s.key)

        ws.cell(row=1, column=1, value=s.title).font = Font(bold=True, size=14, color=INK)
        r = 2
        if s.note:
            c = ws.cell(row=2, column=1, value=s.note)
            c.font = Font(size=9, italic=True, color=MUTED)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(
                start_row=2, start_column=1, end_row=2, end_column=max(2, len(s.columns))
            )
            ws.row_dimensions[2].height = 26
            r = 3

        head = r + 1
        for i, col in enumerate(s.columns, start=1):
            c = ws.cell(row=head, column=i, value=col)
            c.font = Font(bold=True, color=INK)
            c.fill = PatternFill("solid", fgColor=BRAND_SOFT)
            c.border = box
            c.alignment = Alignment(horizontal="right" if _is_money(s, i - 1) else "left")

        for j, row in enumerate(s.rows):
            for i, v in enumerate(row, start=1):
                c = ws.cell(row=head + 1 + j, column=i)
                # NUMBERS GO IN AS NUMBERS. Writing "£1,234.00" as text gives a
                # spreadsheet you cannot sum, which is most of why a person
                # wanted Excel rather than the PDF.
                if isinstance(v, Decimal):
                    c.value = float(v)
                    c.number_format = money_fmt if _is_money(s, i - 1) else "#,##0.00"
                else:
                    c.value = "—" if v is None else str(v)
                c.border = box
                if j % 2:
                    c.fill = PatternFill("solid", fgColor=STRIPE)

        if s.total:
            tr = head + 1 + len(s.rows)
            for i, v in enumerate(s.total, start=1):
                c = ws.cell(row=tr, column=i)
                if isinstance(v, Decimal):
                    c.value = float(v)
                    c.number_format = money_fmt if _is_money(s, i - 1) else "#,##0.00"
                else:
                    c.value = "—" if v is None else str(v)
                c.font = Font(bold=True, color=INK)
                c.fill = PatternFill("solid", fgColor=TOTAL_BG)
                c.border = box

        for i, col in enumerate(s.columns, start=1):
            widest = max([len(str(col))] + [len(str(row[i - 1])) for row in s.rows] or [10])
            ws.column_dimensions[get_column_letter(i)].width = min(46, max(12, widest + 4))
        # The header row stays put while you scroll a long table.
        ws.freeze_panes = ws.cell(row=head + 1, column=1)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ══ PDF ═══════════════════════════════════════════════════════════════════


def to_pdf(report: dict) -> bytes:
    """The document you email to an accountant."""
    from fpdf.enums import XPos, YPos

    from app.core.pdf import branded_pdf, footer, ps

    cur = report["currency"]
    # Landscape: the platform table is five columns wide and portrait squeezes
    # the money into two-line wraps.
    pdf = branded_pdf(_title(report), _period(report), orientation="L")
    usable = pdf.w - 28

    for s in report["sections"]:
        if pdf.get_y() > pdf.h - 60:
            pdf.add_page()

        pdf.set_y(pdf.get_y() + 3)
        pdf.set_x(14)
        pdf.set_fill_color(*_RGB["BRAND"])
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(usable, 9, text=ps("  " + s.title), fill=True,
                 new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_text_color(*_RGB["INK"])

        if s.note:
            pdf.set_x(14)
            pdf.set_font("Helvetica", "I", 8)
            pdf.set_text_color(*_RGB["MUTED"])
            pdf.multi_cell(usable, 4.2, text=ps(s.note))
            pdf.set_text_color(*_RGB["INK"])

        n = len(s.columns)
        # The first column carries names and needs the room; the rest are money.
        first = usable * (0.34 if n > 2 else 0.5)
        rest = (usable - first) / max(1, n - 1)
        widths = [first] + [rest] * (n - 1)

        pdf.set_x(14)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_fill_color(*_RGB["BRAND_SOFT"])
        for i, col in enumerate(s.columns):
            pdf.cell(widths[i], 7, text=ps(col), fill=True,
                     align="R" if _is_money(s, i) else "L", border=0)
        pdf.ln(7)

        pdf.set_font("Helvetica", "", 9)
        for j, row in enumerate(s.rows):
            if pdf.get_y() > pdf.h - 28:
                pdf.add_page()
            pdf.set_x(14)
            striped = j % 2 == 1
            if striped:
                pdf.set_fill_color(*_RGB["STRIPE"])
            for i, v in enumerate(row):
                pdf.cell(
                    widths[i], 6.4,
                    text=ps(fmt(v, money=_is_money(s, i), currency=cur)),
                    align="R" if _is_money(s, i) else "L",
                    fill=striped,
                )
            pdf.ln(6.4)

        if s.total:
            pdf.set_x(14)
            pdf.set_font("Helvetica", "B", 9)
            pdf.set_fill_color(*_RGB["TOTAL_BG"])
            for i, v in enumerate(s.total):
                pdf.cell(
                    widths[i], 7,
                    text=ps(fmt(v, money=_is_money(s, i), currency=cur)),
                    align="R" if _is_money(s, i) else "L",
                    fill=True,
                )
            pdf.ln(9)

    footer(pdf)
    return bytes(pdf.output())


# ══ Word ══════════════════════════════════════════════════════════════════


def to_docx(report: dict) -> bytes:
    """The one somebody will edit before sending it on.

    That is the whole reason Word is on the list: a PDF is final and a
    spreadsheet is data, but an accountant wants to add a paragraph.
    """
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    cur = report["currency"]
    doc = Document()

    def shade(cell, hex_colour: str) -> None:
        """Cell shading has no python-docx API; it is raw OOXML or nothing."""
        el = OxmlElement("w:shd")
        el.set(qn("w:val"), "clear")
        el.set(qn("w:fill"), hex_colour)
        cell._tc.get_or_add_tcPr().append(el)

    h = doc.add_heading(_title(report), level=0)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x3D, 0x3E)
    sub = doc.add_paragraph(_period(report))
    sub.runs[0].italic = True
    sub.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    for s in report["sections"]:
        sh = doc.add_heading(s.title, level=1)
        for run in sh.runs:
            run.font.color.rgb = RGBColor(0x0E, 0x7C, 0x7B)
        if s.note:
            p = doc.add_paragraph(s.note)
            p.runs[0].italic = True
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        body = list(s.rows) + ([s.total] if s.total else [])
        table = doc.add_table(rows=1 + len(body), cols=len(s.columns))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, col in enumerate(s.columns):
            cell = table.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(col)
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x0F, 0x3D, 0x3E)
            shade(cell, BRAND_SOFT)
            if _is_money(s, i):
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        for j, row in enumerate(body):
            is_total = s.total is not None and j == len(body) - 1
            for i, v in enumerate(row):
                cell = table.rows[1 + j].cells[i]
                cell.text = ""
                run = cell.paragraphs[0].add_run(
                    fmt(v, money=_is_money(s, i), currency=cur)
                )
                run.font.size = Pt(9)
                run.bold = is_total
                if is_total:
                    shade(cell, TOTAL_BG)
                elif j % 2:
                    shade(cell, STRIPE)
                if _is_money(s, i):
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    foot = doc.add_paragraph("Generated by DineAI — every plate, every penny")
    foot.runs[0].font.size = Pt(8)
    foot.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
